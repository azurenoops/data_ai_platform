#!/usr/bin/env python3
"""Generate extensive demo data assets for MCP ingestion demos.

Outputs:
- demo-data/landing/manual/*.docx
- demo-data/landing/manual/*.xlsx
- demo-data/sql/*.sql
- demo-data/sql/exports/*.csv
"""

from __future__ import annotations

import csv
import random
from datetime import datetime, timedelta
from pathlib import Path

from docx import Document
from faker import Faker
from openpyxl import Workbook
from openpyxl.styles import Font

fake = Faker()
random.seed(42)
Faker.seed(42)

ROOT = Path(__file__).resolve().parent
LANDING_MANUAL = ROOT / "landing" / "manual"
SQL_DIR = ROOT / "sql"
SQL_EXPORTS = SQL_DIR / "exports"

for path in (LANDING_MANUAL, SQL_DIR, SQL_EXPORTS):
    path.mkdir(parents=True, exist_ok=True)


def make_word_docs() -> None:
    # Operations report doc
    report = Document()
    report.add_heading("Fleet Sustainment Operations Quarterly Report", level=1)
    report.add_paragraph(
        "This report summarizes readiness, maintenance completion, supply health, "
        "mission utilization, and risk mitigation actions across all squadrons."
    )

    sections = [
        "Executive Summary",
        "Readiness Metrics",
        "Maintenance Backlog",
        "Supply Chain Constraints",
        "Personnel and Training",
        "Risk Register",
        "Action Plan",
    ]

    for sec in sections:
        report.add_heading(sec, level=2)
        for _ in range(6):
            report.add_paragraph(fake.paragraph(nb_sentences=7))

    report.add_heading("Key Findings", level=2)
    findings = [
        "Average mission capable rate improved from 71.4% to 78.8% over 90 days.",
        "Top recurring fault family: Hydraulic pressure sensors (18.3% of delayed work orders).",
        "Mean time to repair for avionics failures decreased by 11.2 hours after bench refresh.",
        "Supplier lead time risk remains elevated for high-voltage inverter modules.",
        "Training throughput reached 96 certifications this quarter with 92% first-pass rate.",
    ]
    for item in findings:
        report.add_paragraph(item, style="List Bullet")

    report.add_heading("Appendix A: Program Milestones", level=2)
    for i in range(1, 31):
        report.add_paragraph(
            f"Milestone {i:02d}: {fake.sentence(nb_words=10)} "
            f"Target={fake.date_between(start_date='-180d', end_date='+120d')}"
        )

    report.save(LANDING_MANUAL / "fleet_ops_quarterly_report.docx")

    # Lessons learned / incident review doc
    lessons = Document()
    lessons.add_heading("Operational Lessons Learned and Incident Review", level=1)
    lessons.add_paragraph(
        "This document captures near-miss events, root causes, corrective actions, "
        "verification status, and cross-team recommendations for future readiness cycles."
    )

    for idx in range(1, 41):
        lessons.add_heading(f"Case {idx:02d}: {fake.catch_phrase()}", level=2)
        lessons.add_paragraph("Incident Narrative:", style="Intense Quote")
        lessons.add_paragraph(fake.paragraph(nb_sentences=6))
        lessons.add_paragraph("Root Cause:", style="Intense Quote")
        lessons.add_paragraph(fake.sentence(nb_words=16))
        lessons.add_paragraph("Corrective Actions:", style="Intense Quote")
        for _ in range(3):
            lessons.add_paragraph(fake.sentence(nb_words=14), style="List Number")

    lessons.save(LANDING_MANUAL / "incident_lessons_learned.docx")


def make_excel_workbooks() -> None:
    # Program portfolio workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Portfolio"
    headers = [
        "ProgramId",
        "ProgramName",
        "Directorate",
        "FY",
        "BudgetUSD",
        "ObligatedUSD",
        "BurnRatePct",
        "RiskLevel",
        "MilestonesOnTrack",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    directorates = ["Air", "Surface", "Cyber", "Logistics", "R&D"]
    risks = ["Low", "Medium", "High", "Critical"]

    for i in range(1, 601):
        budget = random.randint(2_000_000, 95_000_000)
        obligated = int(budget * random.uniform(0.35, 0.98))
        burn = round((obligated / budget) * 100, 2)
        ws.append(
            [
                f"PRG-{i:04d}",
                fake.bs().title(),
                random.choice(directorates),
                random.choice([2025, 2026, 2027]),
                budget,
                obligated,
                burn,
                random.choice(risks),
                random.choice(["Yes", "No"]),
            ]
        )

    ws2 = wb.create_sheet("SupplierKPI")
    headers2 = [
        "SupplierId",
        "SupplierName",
        "Category",
        "OnTimeDeliveryPct",
        "DefectRatePct",
        "AvgLeadDays",
        "OpenNCRs",
        "Criticality",
    ]
    ws2.append(headers2)
    for cell in ws2[1]:
        cell.font = Font(bold=True)

    categories = ["Electronics", "Mechanical", "Composite", "Software", "Fasteners"]
    criticality = ["Tier-1", "Tier-2", "Tier-3"]
    for i in range(1, 351):
        ws2.append(
            [
                f"SUP-{i:04d}",
                f"{fake.company()}",
                random.choice(categories),
                round(random.uniform(71.0, 99.7), 2),
                round(random.uniform(0.1, 7.4), 2),
                random.randint(4, 97),
                random.randint(0, 24),
                random.choice(criticality),
            ]
        )

    wb.save(LANDING_MANUAL / "program_portfolio_kpis.xlsx")

    # Mission events workbook
    wb2 = Workbook()
    ev = wb2.active
    ev.title = "MissionEvents"
    headers3 = [
        "EventId",
        "MissionType",
        "Region",
        "LaunchTimeUtc",
        "DurationHours",
        "Status",
        "PrimarySystem",
        "AnomalyFlag",
        "Summary",
    ]
    ev.append(headers3)
    for cell in ev[1]:
        cell.font = Font(bold=True)

    mission_types = ["ISR", "Logistics", "Patrol", "Test", "Training"]
    regions = ["PAC", "ATL", "MED", "CONUS", "EUCOM"]
    statuses = ["Completed", "Aborted", "Delayed", "In Progress"]

    start = datetime.utcnow() - timedelta(days=180)
    for i in range(1, 2201):
        launch = start + timedelta(hours=random.randint(1, 4300))
        ev.append(
            [
                f"EVT-{i:05d}",
                random.choice(mission_types),
                random.choice(regions),
                launch.strftime("%Y-%m-%d %H:%M:%S"),
                round(random.uniform(0.8, 19.2), 2),
                random.choice(statuses),
                random.choice(["Radar", "FlightControl", "Hydraulics", "NavSuite", "PowerBus"]),
                random.choice(["Y", "N"]),
                fake.sentence(nb_words=14),
            ]
        )

    wb2.save(LANDING_MANUAL / "mission_events_analytics.xlsx")


def make_sql_seed_and_exports() -> None:
    sql_file = SQL_DIR / "seed_demo.sql"

    customers = []
    projects = []
    work_orders = []
    findings = []
    telemetry = []

    for i in range(1, 301):
        customers.append((
            i,
            fake.company().replace("'", "''"),
            random.choice(["Defense", "Aerospace", "Maritime", "R&D", "Cyber"]),
            random.choice(["CONUS", "EU", "PAC", "MENA"]),
            fake.date_between(start_date="-4y", end_date="-30d").isoformat(),
        ))

    for i in range(1, 1201):
        c = random.randint(1, 300)
        budget = random.randint(400_000, 40_000_000)
        spent = int(budget * random.uniform(0.2, 1.03))
        projects.append((
            i,
            c,
            fake.catch_phrase().replace("'", "''"),
            random.choice(["Planned", "Active", "At Risk", "Completed", "Paused"]),
            budget,
            spent,
            fake.date_between(start_date="-2y", end_date="-20d").isoformat(),
            fake.date_between(start_date="+20d", end_date="+2y").isoformat(),
        ))

    for i in range(1, 5001):
        p = random.randint(1, 1200)
        severity = random.choice(["Low", "Medium", "High", "Critical"])
        work_orders.append((
            i,
            p,
            random.choice(["Open", "In Progress", "Blocked", "Closed"]),
            severity,
            random.choice(["Avionics", "Hydraulics", "Power", "Hull", "Software", "Comms"]),
            random.randint(1, 120),
            fake.date_between(start_date="-365d", end_date="today").isoformat(),
        ))

    for i in range(1, 9001):
        wo = random.randint(1, 5000)
        findings.append((
            i,
            wo,
            random.choice(["Inspection", "Test", "Audit", "FailureAnalysis"]),
            random.choice(["Pass", "Fail", "Observe"]),
            fake.sentence(nb_words=18).replace("'", "''"),
            fake.date_time_between(start_date="-300d", end_date="now").strftime("%Y-%m-%d %H:%M:%S"),
        ))

    for i in range(1, 20001):
        pid = random.randint(1, 1200)
        ts = fake.date_time_between(start_date="-120d", end_date="now")
        telemetry.append((
            i,
            pid,
            ts.strftime("%Y-%m-%d %H:%M:%S"),
            random.choice(["temperature", "pressure", "voltage", "latency", "throughput"]),
            round(random.uniform(0.01, 999.99), 3),
            random.choice(["nominal", "warning", "critical"]),
        ))

    with sql_file.open("w", encoding="utf-8") as f:
        f.write("SET NOCOUNT ON;\n")
        f.write("IF OBJECT_ID('dbo.TelemetryReadings','U') IS NOT NULL DROP TABLE dbo.TelemetryReadings;\n")
        f.write("IF OBJECT_ID('dbo.InspectionFindings','U') IS NOT NULL DROP TABLE dbo.InspectionFindings;\n")
        f.write("IF OBJECT_ID('dbo.WorkOrders','U') IS NOT NULL DROP TABLE dbo.WorkOrders;\n")
        f.write("IF OBJECT_ID('dbo.Projects','U') IS NOT NULL DROP TABLE dbo.Projects;\n")
        f.write("IF OBJECT_ID('dbo.Customers','U') IS NOT NULL DROP TABLE dbo.Customers;\n")

        f.write("CREATE TABLE dbo.Customers (CustomerId INT PRIMARY KEY, CustomerName NVARCHAR(200), Sector NVARCHAR(80), Region NVARCHAR(40), OnboardedDate DATE);\n")
        f.write("CREATE TABLE dbo.Projects (ProjectId INT PRIMARY KEY, CustomerId INT NOT NULL, ProjectName NVARCHAR(220), Status NVARCHAR(40), BudgetUsd BIGINT, SpentUsd BIGINT, StartDate DATE, TargetEndDate DATE, FOREIGN KEY (CustomerId) REFERENCES dbo.Customers(CustomerId));\n")
        f.write("CREATE TABLE dbo.WorkOrders (WorkOrderId INT PRIMARY KEY, ProjectId INT NOT NULL, Status NVARCHAR(40), Severity NVARCHAR(20), SystemArea NVARCHAR(80), HoursOpen INT, OpenedDate DATE, FOREIGN KEY (ProjectId) REFERENCES dbo.Projects(ProjectId));\n")
        f.write("CREATE TABLE dbo.InspectionFindings (FindingId INT PRIMARY KEY, WorkOrderId INT NOT NULL, FindingType NVARCHAR(50), Result NVARCHAR(20), Notes NVARCHAR(500), CapturedAt DATETIME2, FOREIGN KEY (WorkOrderId) REFERENCES dbo.WorkOrders(WorkOrderId));\n")
        f.write("CREATE TABLE dbo.TelemetryReadings (ReadingId INT PRIMARY KEY, ProjectId INT NOT NULL, CapturedAt DATETIME2, MetricName NVARCHAR(80), MetricValue DECIMAL(12,3), HealthState NVARCHAR(20), FOREIGN KEY (ProjectId) REFERENCES dbo.Projects(ProjectId));\n")

        def write_inserts(table: str, cols: str, rows: list[tuple], batch: int = 400) -> None:
            for i in range(0, len(rows), batch):
                chunk = rows[i : i + batch]
                values = []
                for row in chunk:
                    quoted = []
                    for v in row:
                        if isinstance(v, str):
                            quoted.append(f"'{v}'")
                        else:
                            quoted.append(str(v))
                    values.append(f"({','.join(quoted)})")
                f.write(f"INSERT INTO {table} ({cols}) VALUES\n" + ",\n".join(values) + ";\n")

        write_inserts("dbo.Customers", "CustomerId,CustomerName,Sector,Region,OnboardedDate", customers)
        write_inserts("dbo.Projects", "ProjectId,CustomerId,ProjectName,Status,BudgetUsd,SpentUsd,StartDate,TargetEndDate", projects)
        write_inserts("dbo.WorkOrders", "WorkOrderId,ProjectId,Status,Severity,SystemArea,HoursOpen,OpenedDate", work_orders)
        write_inserts("dbo.InspectionFindings", "FindingId,WorkOrderId,FindingType,Result,Notes,CapturedAt", findings)
        write_inserts("dbo.TelemetryReadings", "ReadingId,ProjectId,CapturedAt,MetricName,MetricValue,HealthState", telemetry)

    # CSV exports for immediate blob ingestion demo
    def write_csv(path: Path, headers: list[str], rows: list[tuple]) -> None:
        with path.open("w", newline="", encoding="utf-8") as fp:
            writer = csv.writer(fp)
            writer.writerow(headers)
            writer.writerows(rows)

    write_csv(SQL_EXPORTS / "customers.csv", ["CustomerId", "CustomerName", "Sector", "Region", "OnboardedDate"], customers)
    write_csv(SQL_EXPORTS / "projects.csv", ["ProjectId", "CustomerId", "ProjectName", "Status", "BudgetUsd", "SpentUsd", "StartDate", "TargetEndDate"], projects)
    write_csv(SQL_EXPORTS / "workorders.csv", ["WorkOrderId", "ProjectId", "Status", "Severity", "SystemArea", "HoursOpen", "OpenedDate"], work_orders)
    write_csv(SQL_EXPORTS / "inspection_findings.csv", ["FindingId", "WorkOrderId", "FindingType", "Result", "Notes", "CapturedAt"], findings)
    write_csv(SQL_EXPORTS / "telemetry_readings.csv", ["ReadingId", "ProjectId", "CapturedAt", "MetricName", "MetricValue", "HealthState"], telemetry)


if __name__ == "__main__":
    make_word_docs()
    make_excel_workbooks()
    make_sql_seed_and_exports()
    print("Demo assets generated under demo-data/")
